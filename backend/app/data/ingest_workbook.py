"""Read-only source ingestion. Never overwrites the supplied workbook."""
from pathlib import Path
import csv
import hashlib
import json
import openpyxl

ROOT = Path(__file__).resolve().parents[3]
SOURCE = ROOT / 'data/source/Hybrid_Physics_ML_Impulse_Generator_Optimiser.xlsx'

def ingest():
    formulas = openpyxl.load_workbook(SOURCE, data_only=False)
    values = openpyxl.load_workbook(SOURCE, data_only=True)
    out = ROOT / 'data/processed'
    out.mkdir(parents=True, exist_ok=True)
    sheets = {}
    for sheet in formulas:
        cells = {c.coordinate: {'value': values[sheet.title][c.coordinate].value,
                                'formula': c.value if c.data_type == 'f' else None}
                 for row in sheet for c in row if c.value is not None}
        sheets[sheet.title] = {'rows': sheet.max_row, 'columns': sheet.max_column, 'cells': cells}
        name = sheet.title.lower().replace(' ', '_')
        with (out / f'{name}.csv').open('w', newline='') as stream:
            csv.writer(stream).writerows(values[sheet.title].values)
    snapshot = {'version': '1.0.0', 'source': SOURCE.name,
                'sha256': hashlib.sha256(SOURCE.read_bytes()).hexdigest(), 'sheets': sheets}
    (ROOT / 'artifacts/reference_workbook_snapshot.json').write_text(json.dumps(snapshot, indent=2))
    rows = list(values['Synthetic Dataset'].values)
    assert len(rows) == 2001 and len(rows[0]) == 22
    from collections import Counter
    counts = dict(Counter(r[1] for r in rows[1:]))
    assert counts == {'Train': 1400, 'Validation': 300, 'Hidden Test': 300}, counts
    summary = {'source_sha256': snapshot['sha256'], 'sheets': list(sheets), 'rows': len(rows)-1,
               'columns': list(rows[0]), 'split_counts': counts,
               'impulse_counts': dict(Counter(r[2] for r in rows[1:]))}
    (out / 'manifest.json').write_text(json.dumps(summary, indent=2))
    print(json.dumps(summary, indent=2))
    return summary

if __name__ == '__main__':
    ingest()
